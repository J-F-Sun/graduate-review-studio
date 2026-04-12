from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image

from ..storage import Storage
from ..utils import clean_text, slugify, truncate
from .common import (
    StructureBuilder,
    infer_heading_level,
    looks_like_figure_caption,
    looks_like_table_caption,
    split_inline_preface_heading,
)


def iter_block_items(parent: DocumentObject):
    parent_elm = parent.element.body
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def _style_heading_level(paragraph: Paragraph) -> int | None:
    style_name = (paragraph.style.name if paragraph.style else "").lower()
    match = re.search(r"heading\s*(\d+)", style_name)
    if match:
        return int(match.group(1))
    zh_match = re.search(r"标题\s*(\d+)", style_name)
    if zh_match:
        return int(zh_match.group(1))
    return None


def _extract_paragraph_image_rel_ids(paragraph: Paragraph) -> list[str]:
    xml = paragraph._element.xml
    return list(dict.fromkeys(re.findall(r'r:embed="(rId\d+)"', xml)))


def _safe_image_size(blob: bytes) -> tuple[int | None, int | None]:
    try:
        image = Image.open(BytesIO(blob))
        return image.width, image.height
    except Exception:
        return None, None


def _extract_caption_candidates(paragraphs: list[str], current_text: str) -> str:
    if looks_like_figure_caption(current_text) or looks_like_table_caption(current_text):
        return current_text
    for candidate in reversed(paragraphs[-2:]):
        if looks_like_figure_caption(candidate) or looks_like_table_caption(candidate):
            return candidate
    return current_text


def parse_docx(
    source_path: Path,
    *,
    storage: Storage,
    paper_id: str,
) -> dict:
    document = Document(str(source_path))
    builder = StructureBuilder("docx")
    recent_texts: list[str] = []
    paragraph_index = 0
    table_index = 0
    image_index = 0

    for block_index, block in enumerate(iter_block_items(document)):
        if isinstance(block, Paragraph):
            text = clean_text(block.text)
            image_rel_ids = _extract_paragraph_image_rel_ids(block)
            current_anchor = {
                "kind": "docx_paragraph",
                "index": paragraph_index,
                "block_index": block_index,
            }
            if text:
                recent_texts.append(text)
                inline_heading = split_inline_preface_heading(text)
                if inline_heading:
                    heading_title, body = inline_heading
                    builder.add_heading(heading_title, 1, page_number=None, source_anchor=current_anchor)
                    if body:
                        builder.add_paragraph(body, role="body", page_number=None, source_anchor=current_anchor)
                else:
                    level = infer_heading_level(text, _style_heading_level(block))
                    if level:
                        builder.add_heading(text, level, page_number=None, source_anchor=current_anchor)
                    else:
                        role = "caption" if looks_like_figure_caption(text) or looks_like_table_caption(text) else "body"
                        builder.add_paragraph(text, role=role, page_number=None, source_anchor=current_anchor)
            for rel_id in image_rel_ids:
                image_part = document.part.related_parts.get(rel_id)
                if not image_part:
                    continue
                blob = image_part.blob
                content_type = image_part.content_type.split("/")[-1]
                filename = f"{slugify(rel_id)}.{content_type}"
                relative_path = storage.save_asset(
                    paper_id,
                    category="images",
                    filename=filename,
                    data=blob,
                )
                width, height = _safe_image_size(blob)
                caption = _extract_caption_candidates(recent_texts, text if text else "文档图片")
                image_anchor = {
                    "kind": "docx_image",
                    "index": image_index,
                    "block_index": block_index,
                    "anchor_paragraph_index": paragraph_index,
                    "relationship_id": rel_id,
                }
                builder.add_image(
                    caption=caption,
                    page_number=None,
                    relative_path=relative_path,
                    nearby_text=truncate(" ".join(recent_texts[-3:]), 260),
                    width=width,
                    height=height,
                    analysis=f"图片来源于 Word 文档，当前根据标题和邻近正文预估其内容为：{truncate(caption or '未命名图片', 120)}。",
                    analysis_status="context-only",
                    source_anchor=image_anchor,
                )
                image_index += 1
            paragraph_index += 1
        elif isinstance(block, Table):
            rows = []
            for row in block.rows:
                rows.append([clean_text(cell.text) for cell in row.cells])
            caption = ""
            for candidate in reversed(recent_texts[-2:]):
                if looks_like_table_caption(candidate):
                    caption = candidate
                    break
            builder.add_table(
                rows,
                caption=caption or "未命名表格",
                page_number=None,
                source_anchor={
                    "kind": "docx_table",
                    "index": table_index,
                    "block_index": block_index,
                    "anchor_paragraph_index": paragraph_index - 1 if paragraph_index > 0 else None,
                },
            )
            table_index += 1

    parsed = builder.finalize()
    parsed["page_count"] = None
    parsed["parser_notes"] = [
        "DOCX 文件按标题样式和编号规则推断章节层级。",
        "DOCX 页码通常不稳定，当前版本不对 Word 文件强行推算页码。",
    ]
    return parsed
