from __future__ import annotations

from collections import defaultdict
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from .utils import clean_text, truncate
from .word_comments import CommentSpec, inject_comments_into_docx


class ExportError(RuntimeError):
    pass


SEVERITY_ORDER = {"严重": 0, "重要": 1, "一般": 2}
SEVERITY_LABEL = {"严重": "严重问题", "重要": "需要重点修改", "一般": "建议优化"}
SEVERITY_HIGHLIGHT = {"严重": "yellow", "重要": "green", "一般": "lightGray"}


def _export_flags(settings: dict[str, Any]) -> dict[str, bool]:
    return settings.get("export", {})


def _sorted_issues(review: dict[str, Any]) -> list[dict[str, Any]]:
    issues = review.get("issues", [])
    return sorted(
        issues,
        key=lambda item: (
            SEVERITY_ORDER.get(item.get("severity", "一般"), 99),
            item.get("location_label", ""),
            item.get("title", ""),
        ),
    )


def build_export_context(
    metadata: dict[str, Any],
    parsed: dict[str, Any],
    review: dict[str, Any],
    settings: dict[str, Any],
) -> dict[str, Any]:
    export_settings = _export_flags(settings)
    issues = _sorted_issues(review)
    return {
        "title": metadata.get("title", "未命名论文"),
        "metadata": metadata,
        "parsed": parsed,
        "review": review,
        "export": export_settings,
        "summary": review.get("summary", "暂无内容概括。"),
        "general_assessment": review.get("general_assessment", ""),
        "dimension_scores": review.get("dimension_scores", []),
        "total_score": review.get("total_score"),
        "pass": review.get("pass"),
        "chapter_reviews": review.get("chapter_reviews", []),
        "issues": issues,
        "rule_hits": review.get("rule_hits", []),
        "innovation": review.get("innovation"),
        "basic_info": [
            ("论文类型", metadata.get("degree_type", "-")),
            ("文件类型", metadata.get("source_type", "-")),
            ("页数", metadata.get("page_count") or parsed.get("page_count") or "-"),
            ("审稿模式", review.get("mode") or metadata.get("review_mode", "-")),
            ("生成来源", review.get("provider", "-")),
        ],
    }


def _require_exportable_review(metadata: dict[str, Any], review: dict[str, Any]) -> None:
    if metadata.get("review_status") != "done":
        raise ExportError("当前论文审稿尚未完成，暂时不能导出。")
    if not isinstance(review, dict) or not review.get("generated_at"):
        raise ExportError("当前论文还没有有效的审稿结果，无法导出。")
    if review.get("generation_mode") in {"failed", "terminated"}:
        raise ExportError(review.get("error_message") or "当前审稿结果无效，无法导出。")


def _issue_lines(issue: dict[str, Any], include_evidence: bool, include_polish: bool) -> list[str]:
    lines = [
        f"- 严重程度：{issue.get('severity', '-')}",
        f"- 问题类型：{issue.get('type', '-')}",
        f"- 位置：{issue.get('location_label', '-')}",
    ]
    if include_evidence and issue.get("evidence"):
        lines.append(f"- 原文：{issue['evidence']}")
    lines.append(f"- 说明：{issue.get('analysis', '-')}")
    lines.append(f"- 修改建议：{issue.get('suggestion', '-')}")
    if include_polish and issue.get("polish_suggestion"):
        lines.append(f"- 润色建议：{issue['polish_suggestion']}")
    return lines


def _grouped_issues_by_title(issues: list[dict[str, Any]]) -> list[tuple[str, list[dict[str, Any]]]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    order: list[str] = []
    for issue in issues:
        key = clean_text(issue.get("title", "")) or clean_text(issue.get("type", "")) or "未命名问题"
        if key not in grouped:
            order.append(key)
        grouped[key].append(issue)
    return [
        (
            key,
            sorted(
                grouped[key],
                key=lambda item: (
                    SEVERITY_ORDER.get(item.get("severity", "一般"), 99),
                    item.get("location_label", ""),
                    item.get("type", ""),
                ),
            ),
        )
        for key in order
    ]


def _severity_bucket(issues: list[dict[str, Any]], severity: str) -> list[dict[str, Any]]:
    return [issue for issue in issues if issue.get("severity", "一般") == severity]


def _language_issues(issues: list[dict[str, Any]]) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for issue in issues:
        issue_type = issue.get("type", "") or ""
        if "语言" in issue_type or issue.get("polish_suggestion"):
            result.append(issue)
    return result


def _priority_summary_line(issues: list[dict[str, Any]]) -> str:
    severe_count = len(_severity_bucket(issues, "严重"))
    important_count = len(_severity_bucket(issues, "重要"))
    general_count = len(_severity_bucket(issues, "一般"))
    return f"当前共识别到 {len(issues)} 个问题，其中严重 {severe_count} 个、重要 {important_count} 个、一般 {general_count} 个。"


def _revision_order_advice(issues: list[dict[str, Any]]) -> str:
    if _severity_bucket(issues, "严重"):
        return "建议先补结构、实验或论证支撑，再处理章节内容完整性，最后统一做语言润色和格式细化。"
    if len(_severity_bucket(issues, "重要")) >= 2:
        return "建议先把章节内容补充完整并理顺论述逻辑，再回头统一修改标题、表达和格式。"
    if _language_issues(issues):
        return "当前以表述和细节问题为主，建议先逐段改句子，再通读一遍检查上下文衔接。"
    return "建议先按问题清单逐项修改，再完整通读一遍，确认相邻段落之间的衔接自然。"


def _has_paragraph_style(document: Any, style_name: str) -> bool:
    try:
        document.styles[style_name]
        return True
    except Exception:
        return False


def _add_paragraph_safe(document: Any, text: str, style: str | None = None) -> Any:
    if style and _has_paragraph_style(document, style):
        return document.add_paragraph(text, style=style)
    return document.add_paragraph(text)


def build_markdown_export(
    metadata: dict[str, Any],
    parsed: dict[str, Any],
    review: dict[str, Any],
    settings: dict[str, Any],
) -> str:
    _require_exportable_review(metadata, review)
    context = build_export_context(metadata, parsed, review, settings)
    export_settings = context["export"]
    lines: list[str] = [f"# {context['title']}", ""]

    if export_settings.get("include_basic_info", True):
        lines.extend(["## 基本信息", ""])
        for label, value in context["basic_info"]:
            lines.append(f"- {label}：{value}")
        lines.append("")

    if export_settings.get("include_summary", True):
        lines.extend(["## 内容概括", "", context["summary"], ""])
        if context["general_assessment"]:
            lines.extend(["## 总评", "", context["general_assessment"], ""])

    if export_settings.get("include_dimension_scores", True):
        lines.extend(["## 评分", ""])
        lines.append(f"- 总分：{context.get('total_score', '-')}")
        lines.append(f"- 是否通过：{'是' if context.get('pass') else '否'}")
        for item in context["dimension_scores"]:
            lines.append(f"- {item['name']}：{item['score']} 分")
        lines.append("")

    if export_settings.get("include_chapter_reviews", True):
        lines.extend(["## 逐章评价", ""])
        for chapter in context["chapter_reviews"]:
            lines.append(f"### {chapter['title']}")
            lines.append("")
            lines.append(chapter.get("assessment", ""))
            lines.append("")

    if export_settings.get("include_issue_list", True):
        lines.extend(["## 重点问题", ""])
        for issue_title, grouped_issues in _grouped_issues_by_title(context["issues"]):
            lines.append(f"### {issue_title}")
            lines.append("")
            for issue in grouped_issues:
                lines.append(f"#### {issue.get('location_label', '位置待确认')}")
                lines.append("")
                lines.extend(
                    _issue_lines(
                        issue,
                        include_evidence=export_settings.get("include_evidence", True),
                        include_polish=export_settings.get("include_polish_suggestions", True),
                    )
                )
                lines.append("")

    if export_settings.get("include_rule_hits", True):
        lines.extend(["## 规则命中情况", ""])
        if context["rule_hits"]:
            for item in context["rule_hits"]:
                lines.append(f"- {item}")
        else:
            lines.append("- 未单独记录命中项。")
        lines.append("")

    if export_settings.get("include_innovation", True) and context.get("innovation"):
        innovation = context["innovation"]
        lines.extend(["## 硕士论文创新点评价", ""])
        for point in innovation.get("points", []):
            lines.append(f"- {point}")
        if innovation.get("assessment"):
            lines.append(f"- 综合评价：{innovation['assessment']}")
        lines.append("")

    return "\n".join(lines).strip() + "\n"


def write_pdf_export(
    metadata: dict[str, Any],
    parsed: dict[str, Any],
    review: dict[str, Any],
    settings: dict[str, Any],
    output_path: Path,
) -> Path:
    _require_exportable_review(metadata, review)
    context = build_export_context(metadata, parsed, review, settings)
    export_settings = context["export"]
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.pdfbase.pdfmetrics import registerFont
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    except ImportError as exc:
        raise ExportError("导出 PDF 需要安装 reportlab==4.4.10。") from exc

    registerFont(UnicodeCIDFont("STSong-Light"))
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleCN",
        parent=styles["Title"],
        fontName="STSong-Light",
        fontSize=20,
        leading=26,
        spaceAfter=12,
        textColor=colors.HexColor("#2d241d"),
    )
    heading_style = ParagraphStyle(
        "HeadingCN",
        parent=styles["Heading2"],
        fontName="STSong-Light",
        fontSize=14,
        leading=20,
        spaceBefore=10,
        spaceAfter=6,
        textColor=colors.HexColor("#264653"),
    )
    body_style = ParagraphStyle(
        "BodyCN",
        parent=styles["BodyText"],
        fontName="STSong-Light",
        fontSize=10.5,
        leading=16,
        wordWrap="CJK",
        spaceAfter=5,
    )
    muted_style = ParagraphStyle(
        "MutedCN",
        parent=body_style,
        textColor=colors.HexColor("#6e6154"),
        fontSize=9.5,
        leading=14,
    )

    story: list[Any] = [Paragraph(escape(context["title"]), title_style), Spacer(1, 3 * mm)]

    def add_block(title: str, text: str, style: ParagraphStyle = body_style) -> None:
        story.append(Paragraph(escape(title), heading_style))
        story.append(Paragraph(escape(text).replace("\n", "<br/>"), style))
        story.append(Spacer(1, 2 * mm))

    if export_settings.get("include_basic_info", True):
        story.append(Paragraph("基本信息", heading_style))
        for label, value in context["basic_info"]:
            story.append(Paragraph(escape(f"{label}：{value}"), body_style))
        story.append(Spacer(1, 2 * mm))

    if export_settings.get("include_summary", True):
        add_block("内容概括", context["summary"])
        if context["general_assessment"]:
            add_block("总评", context["general_assessment"])

    if export_settings.get("include_dimension_scores", True):
        story.append(Paragraph("评分", heading_style))
        story.append(Paragraph(escape(f"总分：{context.get('total_score', '-')}"), body_style))
        story.append(Paragraph(escape(f"是否通过：{'是' if context.get('pass') else '否'}"), body_style))
        for item in context["dimension_scores"]:
            story.append(Paragraph(escape(f"{item['name']}：{item['score']} 分"), body_style))
        story.append(Spacer(1, 2 * mm))

    if export_settings.get("include_chapter_reviews", True):
        story.append(Paragraph("逐章评价", heading_style))
        for chapter in context["chapter_reviews"]:
            story.append(Paragraph(escape(chapter["title"]), muted_style))
            story.append(Paragraph(escape(chapter.get("assessment", "")), body_style))
            story.append(Spacer(1, 1.5 * mm))

    if export_settings.get("include_issue_list", True):
        story.append(Paragraph("重点问题", heading_style))
        for index, issue in enumerate(context["issues"], start=1):
            story.append(Paragraph(escape(f"{index}. {issue['title']}"), muted_style))
            for line in _issue_lines(
                issue,
                include_evidence=export_settings.get("include_evidence", True),
                include_polish=export_settings.get("include_polish_suggestions", True),
            ):
                story.append(Paragraph(escape(line), body_style))
            story.append(Spacer(1, 2 * mm))

    if export_settings.get("include_rule_hits", True):
        story.append(Paragraph("规则命中情况", heading_style))
        if context["rule_hits"]:
            for item in context["rule_hits"]:
                story.append(Paragraph(escape(f"- {item}"), body_style))
        else:
            story.append(Paragraph("未单独记录命中项。", body_style))
        story.append(Spacer(1, 2 * mm))

    if export_settings.get("include_innovation", True) and context.get("innovation"):
        innovation = context["innovation"]
        story.append(Paragraph("硕士论文创新点评价", heading_style))
        for point in innovation.get("points", []):
            story.append(Paragraph(escape(f"- {point}"), body_style))
        if innovation.get("assessment"):
            story.append(Paragraph(escape(f"综合评价：{innovation['assessment']}"), body_style))

    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=context["title"],
    )
    doc.build(story)
    return output_path


def _comment_text_for_issue(issue: dict[str, Any], include_evidence: bool, include_polish: bool) -> str:
    severity = issue.get("severity", "一般")
    lines = [
        f"【{severity} / {issue.get('type', '-') or '审稿意见'}】{issue.get('title', '-')}",
        f"位置：{issue.get('location_label', '-')}",
        f"老师判断：{issue.get('analysis', '-')}",
    ]
    if include_evidence and issue.get("evidence"):
        lines.append(f"原文摘录：{issue['evidence']}")
    lines.append(f"建议这样改：{issue.get('suggestion', '-')}")
    if include_polish and issue.get("polish_suggestion"):
        lines.append(f"可直接参考：{issue['polish_suggestion']}")
    return "\n".join(lines)


def _group_comment_texts(issues: list[dict[str, Any]], include_evidence: bool, include_polish: bool) -> str:
    ordered_issues = sorted(
        issues,
        key=lambda item: (
            SEVERITY_ORDER.get(item.get("severity", "一般"), 99),
            item.get("title", ""),
        ),
    )
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for issue in ordered_issues:
        grouped[issue.get("severity", "一般")].append(issue)

    lines: list[str] = []
    worst = _worst_severity(ordered_issues)
    if worst == "严重":
        lines.append("老师批注：这一处有会影响论文成立性的关键问题，建议优先修改。")
    elif len(ordered_issues) == 1:
        lines.append("老师批注：这一处还有提升空间，建议按下面的方向直接修改。")
    else:
        lines.append(f"老师批注：这一处集中出现 {len(ordered_issues)} 个问题，建议按轻重缓急依次处理。")

    line_number = 1
    for severity in ("严重", "重要", "一般"):
        current_group = grouped.get(severity, [])
        if not current_group:
            continue
        lines.append("")
        lines.append(f"{SEVERITY_LABEL[severity]}：")
        for issue in current_group:
            lines.append(f"{line_number}. {_comment_text_for_issue(issue, include_evidence, include_polish)}")
            line_number += 1

    lines.append("")
    if any(issue.get("severity") == "严重" for issue in ordered_issues):
        lines.append("老师建议：先把结构、实验或论证问题补起来，再统一修改语言表达。")
    elif len(ordered_issues) > 1:
        lines.append("老师建议：先把内容逻辑理顺，再统一调整表述和格式。")
    else:
        lines.append("老师建议：改完这一处后，顺带回看前后两段，确认衔接自然。")

    return "\n".join(lines)


def _worst_severity(issues: list[dict[str, Any]]) -> str:
    ordered_issues = sorted(
        issues,
        key=lambda item: SEVERITY_ORDER.get(item.get("severity", "一般"), 99),
    )
    return ordered_issues[0].get("severity", "一般") if ordered_issues else "一般"


def _descendant_ids(section_id: str, child_map: dict[str, list[str]]) -> set[str]:
    collected = {section_id}
    for child_id in child_map.get(section_id, []):
        collected.update(_descendant_ids(child_id, child_map))
    return collected


def _fallback_docx_paragraph_map(parsed: dict[str, Any], document: Any) -> dict[str, int]:
    mapping: dict[str, int] = {}
    used_indexes: set[int] = set()
    document_texts = [clean_text(paragraph.text) for paragraph in document.paragraphs]
    search_start = 0
    for paragraph in parsed.get("paragraphs", []):
        anchor = paragraph.get("source_anchor") or {}
        if anchor.get("kind") == "docx_paragraph" and isinstance(anchor.get("index"), int):
            mapping[paragraph["id"]] = anchor["index"]
            used_indexes.add(anchor["index"])
            search_start = max(search_start, anchor["index"] + 1)
            continue
        text = clean_text(paragraph.get("text", ""))
        if not text:
            continue
        matched_index = None
        for index in range(search_start, len(document_texts)):
            if index in used_indexes:
                continue
            candidate = document_texts[index]
            if candidate == text or (len(text) > 24 and text[:24] in candidate):
                matched_index = index
                break
        if matched_index is None:
            for index, candidate in enumerate(document_texts):
                if index in used_indexes:
                    continue
                if candidate == text or (len(text) > 24 and text[:24] in candidate):
                    matched_index = index
                    break
        if matched_index is not None:
            mapping[paragraph["id"]] = matched_index
            used_indexes.add(matched_index)
            search_start = max(search_start, matched_index + 1)
    return mapping


def _build_original_docx_anchor_maps(parsed: dict[str, Any], document: Any) -> dict[str, dict[str, int]]:
    paragraph_anchor_map = _fallback_docx_paragraph_map(parsed, document)
    child_map: dict[str, list[str]] = defaultdict(list)
    for section in parsed.get("sections", []):
        if section.get("parent_id"):
            child_map[section["parent_id"]].append(section["id"])

    section_anchor_map: dict[str, int] = {}
    for section in parsed.get("sections", []):
        anchor = section.get("source_anchor") or {}
        if anchor.get("kind") == "docx_paragraph" and isinstance(anchor.get("index"), int):
            section_anchor_map[section["id"]] = anchor["index"]

    for section in parsed.get("sections", []):
        if section["id"] in section_anchor_map:
            continue
        descendant_ids = _descendant_ids(section["id"], child_map)
        for paragraph in parsed.get("paragraphs", []):
            if paragraph.get("section_id") in descendant_ids and paragraph["id"] in paragraph_anchor_map:
                section_anchor_map[section["id"]] = paragraph_anchor_map[paragraph["id"]]
                break

    table_anchor_map: dict[str, int] = {}
    for table in parsed.get("tables", []):
        anchor = table.get("source_anchor") or {}
        paragraph_index = anchor.get("anchor_paragraph_index")
        if isinstance(paragraph_index, int):
            table_anchor_map[table["id"]] = paragraph_index
        elif table.get("section_id") in section_anchor_map:
            table_anchor_map[table["id"]] = section_anchor_map[table["section_id"]]

    image_anchor_map: dict[str, int] = {}
    for image in parsed.get("images", []):
        anchor = image.get("source_anchor") or {}
        paragraph_index = anchor.get("anchor_paragraph_index")
        if isinstance(paragraph_index, int):
            image_anchor_map[image["id"]] = paragraph_index
        elif image.get("section_id") in section_anchor_map:
            image_anchor_map[image["id"]] = section_anchor_map[image["section_id"]]

    return {
        "paragraph": paragraph_anchor_map,
        "table": table_anchor_map,
        "image": image_anchor_map,
        "section": section_anchor_map,
    }


def _append_summary_to_docx(document: Any, context: dict[str, Any], unresolved_issues: list[dict[str, Any]]) -> None:
    export_settings = context["export"]
    issues = context["issues"]
    severe_issues = _severity_bucket(issues, "严重")
    important_issues = _severity_bucket(issues, "重要")
    language_issues = _language_issues(issues)
    document.add_page_break()
    document.add_heading("审稿总结", level=1)

    document.add_heading("总体结论", level=2)
    conclusion_parts = [
        f"从当前稿件来看，总分为 {context.get('total_score', '-')} 分，{'达到通过要求' if context.get('pass') else '暂不建议通过'}。",
        _priority_summary_line(issues),
    ]
    if context["general_assessment"]:
        conclusion_parts.append(context["general_assessment"])
    document.add_paragraph(" ".join(part for part in conclusion_parts if part))

    priority_issues = severe_issues + important_issues
    if priority_issues:
        document.add_heading("建议优先修改的内容", level=2)
        for index, issue in enumerate(priority_issues[:8], start=1):
            _add_paragraph_safe(
                document,
                f"{index}. {issue.get('title', '未命名问题')}（{issue.get('location_label', '位置待确认')}）",
                style="List Number",
            )
            document.add_paragraph(f"老师建议：{issue.get('suggestion', '-')}")

    if language_issues and export_settings.get("include_polish_suggestions", True):
        document.add_heading("可直接修改的语言表达问题", level=2)
        for issue in language_issues[:8]:
            _add_paragraph_safe(
                document,
                f"{issue.get('location_label', '位置待确认')}：{issue.get('title', '未命名问题')}",
                style="List Bullet",
            )
            if issue.get("polish_suggestion"):
                document.add_paragraph(f"可参考改写：{issue['polish_suggestion']}")
            else:
                document.add_paragraph(f"修改建议：{issue.get('suggestion', '-')}")

    document.add_heading("修改顺序建议", level=2)
    document.add_paragraph(_revision_order_advice(issues))

    if export_settings.get("include_basic_info", True):
        document.add_heading("基本信息", level=2)
        for label, value in context["basic_info"]:
            document.add_paragraph(f"{label}：{value}")

    if export_settings.get("include_summary", True):
        document.add_heading("内容概括", level=2)
        document.add_paragraph(context["summary"])
        if context["general_assessment"]:
            document.add_heading("总评", level=2)
            document.add_paragraph(context["general_assessment"])

    if export_settings.get("include_dimension_scores", True):
        document.add_heading("评分与结论", level=2)
        document.add_paragraph(f"总分：{context.get('total_score', '-')}")
        document.add_paragraph(f"是否通过：{'是' if context.get('pass') else '否'}")
        for item in context["dimension_scores"]:
            document.add_paragraph(f"{item['name']}：{item['score']} 分")

    if export_settings.get("include_chapter_reviews", True):
        document.add_heading("逐章简评", level=2)
        for chapter in context["chapter_reviews"]:
            document.add_paragraph(f"{chapter['title']}：{chapter.get('assessment', '')}")

    if export_settings.get("include_issue_list", True):
        document.add_heading("详细问题清单", level=2)
        for index, issue in enumerate(context["issues"], start=1):
            document.add_paragraph(f"{index}. {issue.get('title', '未命名问题')}")
            for line in _issue_lines(
                issue,
                include_evidence=export_settings.get("include_evidence", True),
                include_polish=export_settings.get("include_polish_suggestions", True),
            ):
                document.add_paragraph(line)

    if export_settings.get("include_rule_hits", True):
        document.add_heading("规则命中情况", level=2)
        if context["rule_hits"]:
            for item in context["rule_hits"]:
                _add_paragraph_safe(document, item, style="List Bullet")
        else:
            document.add_paragraph("未单独记录命中项。")

    if export_settings.get("include_innovation", True) and context.get("innovation"):
        innovation = context["innovation"]
        document.add_heading("硕士论文创新点评价", level=2)
        for point in innovation.get("points", []):
            _add_paragraph_safe(document, point, style="List Bullet")
        if innovation.get("assessment"):
            document.add_paragraph(f"综合评价：{innovation['assessment']}")

    if unresolved_issues:
        document.add_heading("未能就地批注的问题", level=2)
        document.add_paragraph("下面这些意见没有找到足够稳定的原文锚点，因此统一追加在文末，便于后续人工核对。")
        for index, issue in enumerate(unresolved_issues, start=1):
            document.add_paragraph(f"{index}. {issue.get('title', '未命名问题')}")
            for line in _issue_lines(
                issue,
                include_evidence=export_settings.get("include_evidence", True),
                include_polish=export_settings.get("include_polish_suggestions", True),
            ):
                document.add_paragraph(line)


def _resolve_issue_anchor(issue: dict[str, Any], anchor_maps: dict[str, dict[str, int]], parsed: dict[str, Any]) -> int | None:
    target_ids = issue.get("target_ids", {})
    for paragraph_id in target_ids.get("paragraph_ids", []):
        if paragraph_id in anchor_maps["paragraph"]:
            return anchor_maps["paragraph"][paragraph_id]
    for table_id in target_ids.get("table_ids", []):
        if table_id in anchor_maps["table"]:
            return anchor_maps["table"][table_id]
    for image_id in target_ids.get("image_ids", []):
        if image_id in anchor_maps["image"]:
            return anchor_maps["image"][image_id]
    for section_id in target_ids.get("section_ids", []):
        if section_id in anchor_maps["section"]:
            return anchor_maps["section"][section_id]

    location_label = issue.get("location_label", "")
    for paragraph in parsed.get("paragraphs", []):
        if location_label and location_label == paragraph.get("location_label") and paragraph["id"] in anchor_maps["paragraph"]:
            return anchor_maps["paragraph"][paragraph["id"]]
    return None


def _prepare_comment_specs(
    parsed: dict[str, Any],
    review: dict[str, Any],
    anchor_maps: dict[str, dict[str, int]],
    *,
    include_evidence: bool,
    include_polish: bool,
) -> tuple[list[CommentSpec], list[dict[str, Any]]]:
    grouped: dict[int, list[dict[str, Any]]] = defaultdict(list)
    unresolved: list[dict[str, Any]] = []
    for issue in _sorted_issues(review):
        anchor_index = _resolve_issue_anchor(issue, anchor_maps, parsed)
        if anchor_index is None:
            unresolved.append(issue)
            continue
        grouped[anchor_index].append(issue)

    comment_specs = [
        CommentSpec(
            paragraph_index=index,
            text=_group_comment_texts(
                issues,
                include_evidence=include_evidence,
                include_polish=include_polish,
            ),
            highlight_color=SEVERITY_HIGHLIGHT.get(_worst_severity(issues), "lightGray"),
        )
        for index, issues in sorted(grouped.items())
    ]
    return comment_specs, unresolved


def _rebuild_docx_from_parsed(parsed: dict[str, Any], data_root: Path) -> tuple[Any, dict[str, dict[str, int]]]:
    try:
        from docx import Document
        from docx.shared import Inches
    except ImportError as exc:
        raise ExportError("导出 Word 需要安装 python-docx。") from exc

    document = Document()
    document.add_heading("根据 PDF 解析结果重建的批注版论文", level=0)
    document.add_paragraph("说明：此版本根据 PDF 结构化解析结果重建，用于承载批注意见，排版不承诺与原 PDF 完全一致。")

    section_map = {section["id"]: section for section in parsed.get("sections", [])}
    paragraph_map = {paragraph["id"]: paragraph for paragraph in parsed.get("paragraphs", [])}
    table_map = {table["id"]: table for table in parsed.get("tables", [])}
    image_map = {image["id"]: image for image in parsed.get("images", [])}

    anchor_maps: dict[str, dict[str, int]] = {
        "paragraph": {},
        "table": {},
        "image": {},
        "section": {},
    }

    def add_body_paragraph(text: str, *, style: str | None = None) -> int:
        document.add_paragraph(text, style=style)
        return len(document.paragraphs) - 1

    def render_section(section_id: str) -> None:
        section = section_map[section_id]
        document.add_heading(section["title"], level=min(max(section.get("level", 1), 1), 9))
        anchor_maps["section"][section_id] = len(document.paragraphs) - 1

        for content_id in section.get("content_ids", []):
            if content_id in paragraph_map:
                paragraph = paragraph_map[content_id]
                anchor_maps["paragraph"][content_id] = add_body_paragraph(paragraph.get("text", ""))
            elif content_id in table_map:
                table = table_map[content_id]
                anchor_maps["table"][content_id] = add_body_paragraph(table.get("caption") or table.get("location_label") or "表格")
                rows = table.get("rows", [])
                if rows:
                    max_cols = max(len(row) for row in rows)
                    doc_table = document.add_table(rows=len(rows), cols=max_cols)
                    for row_index, row in enumerate(rows):
                        for col_index in range(max_cols):
                            doc_table.cell(row_index, col_index).text = row[col_index] if col_index < len(row) else ""
                else:
                    add_body_paragraph("表格未抽取到有效内容。")
            elif content_id in image_map:
                image = image_map[content_id]
                anchor_maps["image"][content_id] = add_body_paragraph(image.get("caption") or image.get("location_label") or "图片")
                image_path = data_root / image["asset_relative_path"]
                if image_path.exists():
                    try:
                        document.add_picture(str(image_path), width=Inches(5.8))
                    except Exception:
                        add_body_paragraph("原图无法嵌入，已保留图片说明。")
                add_body_paragraph(f"图片内容理解：{image.get('content_summary', '暂无图片内容总结。')}")
                if image.get("nearby_text"):
                    add_body_paragraph(f"邻近原文：{image['nearby_text']}")

        for child_id in section.get("child_ids", []):
            render_section(child_id)

    root_sections = [section["id"] for section in parsed.get("sections", []) if not section.get("parent_id")]
    for section_id in root_sections:
        render_section(section_id)

    return document, anchor_maps


def write_docx_export(
    metadata: dict[str, Any],
    parsed: dict[str, Any],
    review: dict[str, Any],
    settings: dict[str, Any],
    *,
    source_path: Path,
    data_root: Path,
    output_path: Path,
) -> Path:
    _require_exportable_review(metadata, review)
    context = build_export_context(metadata, parsed, review, settings)
    export_settings = context["export"]
    include_evidence = export_settings.get("include_evidence", True)
    include_polish = export_settings.get("include_polish_suggestions", True)

    try:
        from docx import Document
    except ImportError as exc:
        raise ExportError("导出 Word 需要安装 python-docx。") from exc

    if metadata.get("source_type") == "docx":
        if not source_path.exists():
            raise ExportError("原始 Word 文件不存在，无法生成批注版导出。")
        document = Document(str(source_path))
        anchor_maps = _build_original_docx_anchor_maps(parsed, document)
    else:
        document, anchor_maps = _rebuild_docx_from_parsed(parsed, data_root)

    comment_specs, unresolved_issues = _prepare_comment_specs(
        parsed,
        review,
        anchor_maps,
        include_evidence=include_evidence,
        include_polish=include_polish,
    )

    _append_summary_to_docx(document, context, unresolved_issues)
    document.save(str(output_path))
    inject_comments_into_docx(
        output_path,
        comment_specs,
        author="审稿老师",
        initials="评",
    )
    return output_path
